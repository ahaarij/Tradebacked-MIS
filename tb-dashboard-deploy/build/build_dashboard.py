#!/usr/bin/env python3
"""
Build the Tradebacked MIS dashboard from the MIS workbook.

Usage:
    python build_dashboard.py "Tradebacked_MIS - 3rd Sep.xlsx"
    python build_dashboard.py MIS.xlsx -o Dashboard.html

What it does:
    1. Reads the cached values (not formulas) of the six sheets the dashboard uses.
    2. Converts them to a compact JSON snapshot (dates as YYYY-MM-DD).
    3. Inlines src/dashboard.css, src/dashboard.js and the snapshot into
       src/dashboard.html, producing ONE self-contained HTML file.

Requirements: Python 3.8+, openpyxl  (pip install openpyxl)

Note: values are read as last calculated by Excel, so save the workbook in Excel
before building (a file only edited by a script may have no cached values).
"""
import argparse
import datetime as dt
import json
import math
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "src")

# Sheet name -> number of columns to read (A onwards)
SHEETS = {
    "Transactions Master": 17,  # A:Q  Txn ID ... Products
    "Company Master": 13,       # A:M  Company ID ... Remarks
    "Dashboard Data": 11,       # A:K  segment tables + liquidation sales
    "Payment Receipts": 6,      # A:F  Date, Company, Deal No., Invoice No., Payment, Receipts
    "Custodian Master": 16,     # A:P  custodian fee invoices
    "Pending Deals": 8,         # A:H  pending deals + inspection projection
}
REQUIRED = {"Transactions Master"}


def convert(v):
    """Make a cell value JSON-safe."""
    if v is None:
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return None if math.isnan(v) else round(v, 6)
    if isinstance(v, (int, bool)):
        return v
    return str(v)


def read_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    names = {n.strip().lower(): n for n in wb.sheetnames}
    sheets, missing = {}, []
    for name, ncols in SHEETS.items():
        real = names.get(name.lower())
        if real is None:
            missing.append(name)
            continue
        ws = wb[real]
        grid = [[convert(c) for c in row]
                for row in ws.iter_rows(min_row=1, max_col=ncols, values_only=True)]
        while grid and all(x is None for x in grid[-1]):
            grid.pop()
        sheets[name] = grid
    wb.close()
    if REQUIRED & set(missing):
        sys.exit(f"'{path}' has no 'Transactions Master' sheet - is this the Tradebacked MIS file?")
    for m in missing:
        print(f"warning: sheet '{m}' not found; its section will be empty", file=sys.stderr)
    return sheets


def build(xlsx, out):
    sheets = read_workbook(xlsx)
    payload = {
        "meta": {"source": os.path.basename(xlsx), "extracted": dt.date.today().isoformat()},
        "sheets": sheets,
    }
    data = json.dumps(payload, separators=(",", ":"), ensure_ascii=False).replace("</", "<\\/")

    with open(os.path.join(SRC, "dashboard.html"), encoding="utf-8") as f:
        html = f.read()
    with open(os.path.join(SRC, "dashboard.css"), encoding="utf-8") as f:
        css = f.read()
    with open(os.path.join(SRC, "dashboard.js"), encoding="utf-8") as f:
        js = f.read()

    for marker in ("/*__CSS__*/", "/*__JS__*/", "__DATA__"):
        if marker not in html:
            sys.exit(f"src/dashboard.html is missing the {marker} placeholder")
    html = html.replace("/*__CSS__*/", css).replace("/*__JS__*/", js).replace("__DATA__", data)

    with open(out, "w", encoding="utf-8") as f:
        f.write(html)
    rows = {k: len(v) for k, v in sheets.items()}
    print(f"Built {out} ({len(html) / 1024:.0f} KB) from {os.path.basename(xlsx)}: {rows}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Build the Tradebacked MIS HTML dashboard.")
    ap.add_argument("workbook", help="path to the MIS .xlsx / .xlsm file")
    ap.add_argument("-o", "--out", default="Tradebacked_MIS_Dashboard.html", help="output HTML file")
    a = ap.parse_args()
    build(a.workbook, a.out)
